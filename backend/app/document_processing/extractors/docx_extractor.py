"""DOCX text extraction. Word documents don't carry reliable page-break
metadata without actually rendering/paginating them (which needs a real
layout engine, not a text extractor) — the whole document is treated as
one page. This is a documented simplification, not a bug: `page_count`
for a DOCX will always come out as 1."""
from io import BytesIO

from docx import Document


def extract_docx(data: bytes) -> list[str]:
    document = Document(BytesIO(data))
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip():
                paragraphs.append(row_text)

    return ["\n".join(paragraphs)]
