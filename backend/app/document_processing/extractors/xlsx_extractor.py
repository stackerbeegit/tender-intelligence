"""XLSX text extraction — one "page" per worksheet (e.g. a BOQ's line
items), cells joined with tabs so columns stay visually distinguishable
in the stored text."""
from io import BytesIO

from openpyxl import load_workbook


def extract_xlsx(data: bytes) -> list[str]:
    workbook = load_workbook(BytesIO(data), data_only=True, read_only=True)
    pages: list[str] = []

    for sheet in workbook.worksheets:
        lines = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(cell) if cell is not None else "" for cell in row]
            if any(cell.strip() for cell in cells):
                lines.append("\t".join(cells))
        pages.append(f"[Sheet: {sheet.title}]\n" + "\n".join(lines))

    return pages
