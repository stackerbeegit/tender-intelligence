"""
Downloadable bid pack: the draft as an editable Word document, the
submission checklist as an Excel workbook, and both in one zip. Pure
functions from stored data to bytes.
"""
import io
import re
import zipfile
from datetime import datetime
from typing import Any

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from app.bidding.grounding import PLACEHOLDER
from app.matching.text import format_inr
from app.models.enums import ChecklistItemStatus

STATUS_LABELS = {
    ChecklistItemStatus.AVAILABLE.value: "Available",
    ChecklistItemStatus.TO_PREPARE.value: "To prepare",
    ChecklistItemStatus.MISSING.value: "Missing",
    ChecklistItemStatus.NOT_APPLICABLE.value: "Not applicable",
}


def safe_filename(text: str, fallback: str = "tender") -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "-", text).strip("-")[:60]
    return slug or fallback


def _citation(c: dict[str, Any] | None) -> str:
    return f"{c['file_name']} p.{c['page_number']}" if c else "—"


# ------------------------------------------------------------------ docx


def _add_text(paragraph, text: str, bold: bool = False) -> None:
    """Writes `text` into `paragraph`, highlighting every [TO FILL] span."""
    position = 0
    for match in PLACEHOLDER.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()]).bold = bold
        run = paragraph.add_run(match.group(0))
        run.bold = True
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:]).bold = bold


def _paragraphs(doc, text: str) -> None:
    # Word's paragraph spacing already separates blocks; blank lines would double it.
    for block in text.split("\n"):
        if block.strip():
            _add_text(doc.add_paragraph(), block)


def _table(doc, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = ""
        cell.paragraphs[0].add_run(header).bold = True
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = ""
            _add_text(cell.paragraphs[0], value)


def draft_docx(content: dict[str, Any], version: int) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    tender = content.get("tender", {})

    doc.add_heading(f"Draft bid — {tender.get('title', '')}", level=0)
    banner = doc.add_paragraph()
    run = banner.add_run(content["disclaimer"])
    run.bold = True
    run.font.color.rgb = RGBColor(0xB4, 0x53, 0x09)
    meta = [
        f"Tender number: {tender.get('tender_number') or '—'}",
        f"Issued by: {tender.get('organisation') or '—'}",
        f"Closing: {tender.get('closing_date') or '—'}",
        f"Draft version {version}, generated {content.get('generated_at', '')[:16].replace('T', ' ')} UTC "
        f"by {content.get('generated_by_model', '')}",
    ]
    for line in meta:
        doc.add_paragraph(line).runs[0].font.size = Pt(9)

    doc.add_heading("1. Covering letter", level=1)
    _paragraphs(doc, content["covering_letter"])

    doc.add_heading("2. Company profile", level=1)
    _paragraphs(doc, content["company_overview"])

    doc.add_heading("3. Eligibility compliance", level=1)
    _table(
        doc,
        ["Criterion", "Our compliance", "Tender reference"],
        [[e["criterion"], e["statement"], _citation(e.get("citation"))] for e in content["eligibility_compliance"]],
    )

    doc.add_heading("4. Technical compliance", level=1)
    _table(
        doc,
        ["Requirement", "Our response", "Tender reference"],
        [[t["requirement"], t["response"], _citation(t.get("citation"))] for t in content["technical_compliance"]],
    )

    if content["manpower_plan"]:
        doc.add_heading("5. Manpower deployment", level=1)
        _table(
            doc,
            ["Role", "Count", "Min. experience", "Proposed resource", "Tender reference"],
            [
                [
                    str(m["role"]),
                    str(m["count"] if m["count"] is not None else "—"),
                    f"{m['min_experience_years']:g} years" if m.get("min_experience_years") is not None else "—",
                    m["proposed_resource"],
                    _citation(m.get("citation")),
                ]
                for m in content["manpower_plan"]
            ],
        )

    doc.add_heading("6. Relevant past projects", level=1)
    if content["relevant_projects"]:
        _table(
            doc,
            ["Project", "Client", "Value", "Year"],
            [
                [
                    p["project_name"],
                    f"{p.get('client_name') or '—'} ({p['customer_type']})",
                    format_inr(p["value_inr"]) if p.get("value_inr") is not None else "—",
                    str(p.get("year") or "—"),
                ]
                for p in content["relevant_projects"]
            ],
        )
    else:
        _add_text(doc.add_paragraph(), "[TO FILL: relevant past projects — none matched in the company profile]")

    if content["assumptions"]:
        doc.add_heading("7. Assumptions and points to confirm", level=1)
        for item in content["assumptions"]:
            _add_text(doc.add_paragraph(style="List Bullet"), item)

    doc.add_heading("Enclosures", level=1)
    _table(
        doc,
        ["#", "Document", "Mandatory", "Status"],
        [
            [
                str(i),
                e["title"],
                {True: "Yes", False: "No"}.get(e.get("mandatory"), "Not stated"),
                STATUS_LABELS.get(e.get("status"), e.get("status") or ""),
            ]
            for i, e in enumerate(content["enclosures"], start=1)
        ],
    )

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ------------------------------------------------------------------ xlsx

CHECKLIST_COLUMNS = [
    ("#", 5),
    ("Document", 48),
    ("Kind", 14),
    ("Mandatory", 11),
    ("Format instructions", 32),
    ("Form / annexure", 16),
    ("Tender reference", 22),
    ("Tender clause", 60),
    ("Status", 15),
    ("Owner", 20),
    ("Note", 36),
    ("Why pre-filled", 36),
]


def checklist_xlsx(tender: dict[str, Any], rows: list[dict[str, Any]], generated_at: datetime) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Checklist"
    ws.append([name for name, _ in CHECKLIST_COLUMNS])
    for index, (_, width) in enumerate(CHECKLIST_COLUMNS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=index).column_letter].width = width
        cell = ws.cell(row=1, column=index)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4C1D95")
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for i, row in enumerate(rows, start=1):
        ws.append(
            [
                i,
                row["title"],
                row["kind"].replace("_", " "),
                {True: "Yes", False: "No"}.get(row.get("mandatory"), "Not stated"),
                row.get("format_instructions") or "",
                row.get("form_reference") or "",
                f"{row['source_file_name']} p.{row['source_page']}" if row.get("source_page") else "Added manually",
                row.get("source_clause") or "",
                STATUS_LABELS[row["status"]],
                row.get("owner") or "",
                row.get("note") or "",
                row.get("prefill_reason") or "",
            ]
        )
    widths = [width for _, width in CHECKLIST_COLUMNS]
    for cells in ws.iter_rows(min_row=2):
        lines = 1
        for cell, width in zip(cells, widths):
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            lines = max(lines, -(-len(str(cell.value or "")) // max(width - 2, 1)))
        # Neither Excel nor LibreOffice re-measures wrapped rows in a generated
        # file, so set the height or long clauses overlap the next row.
        ws.row_dimensions[cells[0].row].height = 15 * lines
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"

    status_choices = DataValidation(type="list", formula1='"' + ",".join(STATUS_LABELS.values()) + '"', allow_blank=False)
    ws.add_data_validation(status_choices)
    if rows:
        status_choices.add(f"I2:I{len(rows) + 1}")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:L{max(len(rows), 1) + 1}"

    about = wb.create_sheet("About")
    for key, value in [
        ("Tender", tender.get("title")),
        ("Tender number", tender.get("tender_number")),
        ("Issued by", tender.get("organisation")),
        ("Closing", tender.get("closing_date")),
        ("Generated", generated_at.strftime("%d %b %Y %H:%M UTC")),
        (
            "How to read",
            "Rows come from the tender's 'documents to be submitted' clauses, each with the page it was found on. "
            "Status is pre-filled as Available only where the company profile confirms it; check every row.",
        ),
    ]:
        about.append([key, value or "—"])
    about.column_dimensions["A"].width = 16
    about.column_dimensions["B"].width = 100
    for cell in about["A"]:
        cell.font = Font(bold=True)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ------------------------------------------------------------------- zip


def bid_pack_zip(base_name: str, draft: bytes | None, checklist: bytes) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(f"{base_name}-checklist.xlsx", checklist)
        if draft is not None:
            archive.writestr(f"{base_name}-draft-bid.docx", draft)
        archive.writestr(
            "README.txt",
            "Bid pack generated by Stackerbee Tender Intelligence.\n\n"
            "- draft-bid.docx: AI DRAFT. Review every statement and complete every highlighted [TO FILL] item.\n"
            "- checklist.xlsx: documents the tender asks for, with the page each requirement is on.\n\n"
            "Nothing here has been submitted. Submission is done by the bid team on the tender portal.\n",
        )
    return buffer.getvalue()
