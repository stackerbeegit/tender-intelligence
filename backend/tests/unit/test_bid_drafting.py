"""
The draft-bid guarantees, tested without a database or a model:

- a draft stating an amount, certification, year count or reference that
  isn't in its inputs is rejected (the "no invented credentials" rule);
- text inside [TO FILL] placeholders is not treated as a claim;
- every requirement/criterion gets a row even if the model skipped it;
- the .docx/.xlsx/.zip downloads open and contain what they should.
"""
import copy
import io
import zipfile
from datetime import datetime, timezone

import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from openpyxl import load_workbook

from app.bidding import draft as draft_rules
from app.bidding import export, grounding
from app.bidding.checklist import prefill
from app.matching.types import Capability, CompanyProfile, TurnoverYear
from app.models.enums import ChecklistItemStatus, SubmissionDocumentKind, TenderCategory
from tests.support.fake_llm import load_rfp_responses

NOW = datetime(2026, 10, 22, 9, 0, tzinfo=timezone.utc)
CITE = {"document_id": "doc-1", "file_name": "RFP.pdf", "page_number": 3}


def _fact(ref, kind, summary, clause, details=None, assessment=None):
    return draft_rules.FactRef(
        ref=ref, kind=kind, summary=summary, clause=clause, citation=CITE, details=details or {}, assessment=assessment
    )


def make_inputs() -> draft_rules.DraftInputs:
    return draft_rules.DraftInputs(
        tender={"title": "Integrated Citizen Services Portal", "tender_number": "DIT/PB/2026/ICSP-014",
                "organisation": "Department of IT, Punjab", "closing_date": "30 Nov 2026 15:00"},
        requirements=[
            _fact("R1", "contract", "36-month contract", "The contract period shall be 36 months from the date of Go-Live."),
            _fact("R2", "manpower", "4 Java developers", "Java Developer - 4 resources with minimum 5 years of experience each.",
                  {"role": "Java Developer", "resource_count": 4, "min_experience_years": 5}),
            _fact("R3", "security", "CERT-In audit", "The application must undergo a security audit by a CERT-In empanelled auditor."),
        ],
        eligibility=[
            _fact("E1", "turnover", "Average turnover ₹10 Cr",
                  "The bidder must have an average annual turnover of at least Rs. 10 Crore during the last three financial years.",
                  {"amount_inr": 100000000, "period_years": 3}, "MET"),
            _fact("E2", "certification", "ISO 27001", "The bidder must hold a valid ISO 27001 certification.", {}, "NOT_MET"),
        ],
        company={"company_name": "Stackerbee Technologies", "turnover_by_year": [{"financial_year": "2024-25", "turnover": "₹12 Cr"}],
                 "years_in_business": 9, "certifications": ["ISO 9001"], "capabilities": [], "not_recorded": []},
        projects=[draft_rules.ProjectRef("P1", "State portal", "Govt of Punjab", "government", "software_development",
                                         30000000.0, 2024, ["Java"])],
        enclosures=[{"title": "PAN card", "mandatory": True, "status": "to_prepare", "form_reference": None}],
        extra_amounts={100000000.0, 120000000.0, 30000000.0},
        extra_years={3, 5, 9},
    )


def good_response() -> dict:
    return {
        "covering_letter": "We, Stackerbee Technologies, submit our bid.\n[TO FILL: name of authorised signatory]",
        "company_overview": "Stackerbee Technologies has 9 years in business and turnover of ₹12 Cr in 2024-25.",
        "eligibility_compliance": [
            {"criterion_ref": "E1", "statement": "Average turnover of ₹12 Cr against ₹10 Cr required."},
            {"criterion_ref": "E2", "statement": "Requires ISO 27001. [TO FILL: ISO 27001 certificate — not held per profile]"},
        ],
        "technical_compliance": [{"requirement_ref": "R3", "response": "Complied: CERT-In empanelled auditor."}],
        "manpower_plan": [{"requirement_ref": "R2", "proposed_resource": "[TO FILL: names of 4 developers]"}],
        "relevant_project_refs": ["P1"],
        "assumptions": ["Pricing is in the BOQ."],
    }


def _verify(data):
    inputs = make_inputs()
    return draft_rules.verify(data, inputs, draft_rules.build_prompt(inputs))


# -------------------------------------------------------------- grounding


def test_a_grounded_draft_passes():
    assert _verify(good_response()) == []


@pytest.mark.parametrize(
    "field, text, expected",
    [
        ("company_overview", "Our turnover is ₹50 Cr.", "amount ₹50 Cr"),
        ("company_overview", "Annual revenue of Rs. 2,50,00,000.", "amount ₹2.5 Cr"),
        ("company_overview", "We have 15 years of experience.", "'15 years'"),
        ("company_overview", "We are CMMI Level 5 appraised.", "CMMI Level 5"),
        ("covering_letter", "We hold ISO 20000 certification.", "ISO 20000"),
    ],
)
def test_invented_figures_and_certifications_are_rejected(field, text, expected):
    data = good_response()
    data[field] += " " + text
    problems = _verify(data)
    assert len(problems) == 1 and expected in problems[0]


def test_an_unmet_criterion_cannot_be_presented_as_complete():
    data = good_response()
    data["eligibility_compliance"][1]["statement"] = "We hold a valid ISO 27001 certificate."
    [problem] = _verify(data)
    assert "eligibility E2" in problem and "NOT_MET" in problem


def test_a_certification_the_company_lacks_cannot_be_claimed_elsewhere():
    data = good_response()
    data["company_overview"] += " We are ISO 27001 certified."
    [problem] = _verify(data)
    assert "company overview" in problem and "ISO 27001" in problem


def test_text_inside_placeholders_is_not_a_claim():
    data = good_response()
    data["company_overview"] += " [TO FILL: turnover of ₹99 Cr, CMMI Level 5, 40 years — confirm]"
    assert _verify(data) == []


def test_unknown_references_are_rejected():
    data = good_response()
    data["eligibility_compliance"].append({"criterion_ref": "E9", "statement": "Complied."})
    data["manpower_plan"].append({"requirement_ref": "R1", "proposed_resource": "[TO FILL: x]"})
    data["relevant_project_refs"].append("P7")
    problems = _verify(data)
    assert any("'E9'" in p for p in problems)
    assert any("'R1' is not a manpower requirement" in p for p in problems)
    assert any("'P7'" in p for p in problems)


@pytest.mark.parametrize(
    "text, expected",
    [("₹10 Cr", [1e8]), ("Rs. 5,00,000", [5e5]), ("2 crore", [2e7]), ("INR 1.5 lakh", [1.5e5]), ("no money", [])],
)
def test_amount_parsing(text, expected):
    assert grounding.amounts(text) == expected


# --------------------------------------------------------------- assembly


def test_assembly_gives_every_requirement_a_row_and_lists_placeholders():
    data = good_response()
    data["technical_compliance"] = []  # model skipped R1 and R3
    content, placeholders = draft_rules.assemble(data, make_inputs(), NOW, "fake-model")

    assert [t["ref"] for t in content["technical_compliance"]] == ["R1", "R3"]
    assert all(t["response"].startswith("[TO FILL") for t in content["technical_compliance"])
    assert [m["ref"] for m in content["manpower_plan"]] == ["R2"]
    assert content["manpower_plan"][0]["count"] == 4
    assert content["eligibility_compliance"][0]["citation"] == CITE
    assert content["relevant_projects"][0]["project_name"] == "State portal"  # from the profile, not the model
    assert content["disclaimer"].startswith("AI DRAFT")
    sections = {p["section"] for p in placeholders}
    assert {"technical R1", "technical R3", "manpower R2", "covering letter"} <= sections


def test_the_canned_rfp_draft_is_grounded_against_its_own_fixture_shape():
    """Guards the fixture used by the API tests: it must only use refs
    the real RFP produces."""
    canned = load_rfp_responses()["bid_draft"]
    assert {e["criterion_ref"] for e in canned["eligibility_compliance"]} == {"E1", "E2", "E3", "E4", "E5"}
    assert "R7" not in {t["requirement_ref"] for t in canned["technical_compliance"]}


# ---------------------------------------------------------------- prefill


def _profile(**overrides) -> CompanyProfile:
    base = dict(
        msme_registered=True, startup_registered=False, turnover_by_year=[TurnoverYear("2024-25", 1.2e8)],
        net_worth_inr=None, years_in_business=9, employee_count=120, locations=["Punjab"],
        certifications=["ISO 27001"], capabilities=[Capability(TenderCategory.SOFTWARE_DEVELOPMENT, [], ["CMMI Level 3"])],
        experience=[],
    )
    base.update(overrides)
    return CompanyProfile(**base)


@pytest.mark.parametrize(
    "kind, clause, msme_ref, expected",
    [
        (SubmissionDocumentKind.CERTIFICATE, "Copy of valid ISO 27001 certificate.", None, ChecklistItemStatus.AVAILABLE),
        (SubmissionDocumentKind.CERTIFICATE, "Copy of valid CMMI Level 3 certificate.", None, ChecklistItemStatus.AVAILABLE),
        (SubmissionDocumentKind.CERTIFICATE, "Copy of valid ISO 9001 certificate.", None, ChecklistItemStatus.TO_PREPARE),
        (SubmissionDocumentKind.EMD, "Bank Guarantee towards EMD, or a valid Udyam / NSIC certificate.", "UDYAM-1",
         ChecklistItemStatus.AVAILABLE),
        (SubmissionDocumentKind.EMD, "Bank Guarantee towards EMD, or a valid Udyam / NSIC certificate.", None,
         ChecklistItemStatus.TO_PREPARE),
        (SubmissionDocumentKind.REGISTRATION, "Copy of PAN card and GST registration certificate.", "UDYAM-1",
         ChecklistItemStatus.TO_PREPARE),
    ],
)
def test_checklist_prefill_only_marks_what_the_profile_confirms(kind, clause, msme_ref, expected):
    status, reason = prefill(kind, clause, _profile(), msme_ref, None)
    assert status is expected
    assert (reason is not None) == (expected is ChecklistItemStatus.AVAILABLE)


def test_no_profile_means_nothing_is_prefilled():
    status, _ = prefill(SubmissionDocumentKind.CERTIFICATE, "ISO 27001 certificate", None, None, None)
    assert status is ChecklistItemStatus.TO_PREPARE


# ---------------------------------------------------------------- exports


def test_draft_docx_opens_with_disclaimer_tables_and_highlighted_placeholders():
    content, _ = draft_rules.assemble(good_response(), make_inputs(), NOW, "fake-model")
    doc = Document(io.BytesIO(export.draft_docx(content, version=2)))

    text = "\n".join(p.text for p in doc.paragraphs)
    assert "AI DRAFT — review before use" in text
    assert "Draft version 2" in text
    highlighted = [r.text for p in doc.paragraphs for r in p.runs if r.font.highlight_color == WD_COLOR_INDEX.YELLOW]
    assert "[TO FILL: name of authorised signatory]" in highlighted
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    assert "RFP.pdf p.3" in cells
    assert "State portal" in cells
    assert "PAN card" in cells


def test_checklist_xlsx_has_one_row_per_item_with_a_status_dropdown():
    rows = [
        {"title": "ISO 27001 certificate", "kind": "certificate", "mandatory": True, "format_instructions": None,
         "form_reference": None, "source_clause": "Copy of valid ISO 27001 certificate.", "source_file_name": "RFP.pdf",
         "source_page": 4, "status": "available", "owner": "Asha", "note": None, "prefill_reason": "ISO 27001 recorded"},
        {"title": "Board resolution", "kind": "other", "mandatory": None, "format_instructions": None,
         "form_reference": None, "source_clause": None, "source_file_name": None, "source_page": None,
         "status": "to_prepare", "owner": None, "note": "ask CS", "prefill_reason": None},
    ]
    wb = load_workbook(io.BytesIO(export.checklist_xlsx({"title": "ICSP"}, rows, NOW)))
    ws = wb["Checklist"]
    values = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    assert values[0][1] == "ISO 27001 certificate" and values[0][6] == "RFP.pdf p.4" and values[0][8] == "Available"
    assert values[1][6] == "Added manually" and values[1][3] == "Not stated" and values[1][8] == "To prepare"
    assert ws.data_validations.dataValidation[0].formula1 == '"Available,To prepare,Missing,Not applicable"'
    assert wb["About"]["B1"].value == "ICSP"


def test_bid_pack_zip_contains_both_files_and_a_readme():
    archive = zipfile.ZipFile(io.BytesIO(export.bid_pack_zip("ICSP-014", b"docx-bytes", b"xlsx-bytes")))
    assert sorted(archive.namelist()) == ["ICSP-014-checklist.xlsx", "ICSP-014-draft-bid.docx", "README.txt"]
    assert "Nothing here has been submitted" in archive.read("README.txt").decode()


def test_bid_pack_without_a_draft_still_has_the_checklist():
    archive = zipfile.ZipFile(io.BytesIO(export.bid_pack_zip("x", None, b"xlsx")))
    assert "x-draft-bid.docx" not in archive.namelist()


def test_safe_filename():
    assert export.safe_filename("DIT/PB/2026/ICSP-014") == "DIT-PB-2026-ICSP-014"
    assert export.safe_filename("///") == "tender"


def test_grounding_check_is_pure():
    data = good_response()
    before = copy.deepcopy(data)
    _verify(data)
    assert data == before
