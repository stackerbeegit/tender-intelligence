"""
Sales-triggered bid preparation end to end, on the representative RFP
(tests/fixtures/documents/sample_rfp.pdf): real upload, processing and
analysis (scripted model), then "Interested — prepare bid", the checklist,
the grounded draft and the three downloads.
"""
import copy
import io
import json
import zipfile
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook

from app.core.config import get_settings
from app.core.exceptions import LLMNotConfiguredError
from app.models.audit_log import AuditLog
from tests.support.fake_llm import FakeLLMClient, load_rfp_responses

FIXTURES = Path(__file__).parent.parent / "fixtures"
EXPECTED = json.loads((FIXTURES / "intelligence" / "sample_rfp_expected.json").read_text())
PROFILE = {
    "company_name": "Stackerbee Technologies",
    "msme_registered": True,
    "msme_certificate_ref": "UDYAM-PB-00-0001234",
    "turnover_by_year": [{"financial_year": fy, "amount_inr": 120000000} for fy in ("2024-25", "2023-24", "2022-23")],
    "employee_count": 120,
    "locations": ["Punjab"],
    "certifications": ["ISO 27001"],
}


def _login(client, email, role_name, full_name="Test User"):
    client.post(
        "/api/v1/auth/register",
        json={"email": email, "password": "StrongPass123", "full_name": full_name, "role_name": role_name},
    )
    token = client.post("/api/v1/auth/login", json={"email": email, "password": "StrongPass123"}).json()
    return {"Authorization": f"Bearer {token['access_token']}"}


@pytest.fixture()
def llm(monkeypatch) -> FakeLLMClient:
    fake = FakeLLMClient()
    monkeypatch.setattr("app.services.tender_intelligence_service.get_llm_client", lambda: fake)
    monkeypatch.setattr("app.services.bid_preparation_service.get_llm_client", lambda: fake)
    monkeypatch.setattr(get_settings(), "opportunity_auto_promote_min_score", 0)  # tested separately
    return fake


@pytest.fixture()
def admin(client):
    return _login(client, "admin-b@stackerbee.com", "admin")


@pytest.fixture()
def sales(client):
    return _login(client, "sales-b@stackerbee.com", "sales", "Sales Person")


def _analysed_tender(client, tender, admin, with_profile=True):
    if with_profile:
        assert client.put("/api/v1/profile", headers=admin, json=PROFILE).status_code == 200
    with open(FIXTURES / "documents" / "sample_rfp.pdf", "rb") as f:
        client.post(
            f"/api/v1/tenders/{tender.id}/documents",
            headers=admin,
            files={"file": ("RFP.pdf", f, "application/pdf")},
            data={"document_type": "rfp"},
        )
    analysis = client.post(f"/api/v1/tenders/{tender.id}/analysis", headers=admin).json()
    assert analysis["analysis"]["status"] == "done", analysis["analysis"]["error"]
    return analysis


def test_sales_user_gets_an_opportunity_checklist_and_grounded_draft(
    client, sample_tender, storage_backend, llm, admin, sales, outbox, db_session
):
    _analysed_tender(client, sample_tender, admin)

    response = client.post(f"/api/v1/tenders/{sample_tender.id}/bid-preparation", headers=sales)

    assert response.status_code == 200
    body = response.json()
    assert body["opportunity_created"] is True
    assert body["draft_error"] is None, body["draft_problems"]

    # The opportunity exists, the sales user is on it, and the workflow did NOT move.
    opp = client.get(f"/api/v1/opportunities/{body['opportunity_id']}", headers=sales).json()
    assert opp["status"] == "new"
    assert [(a["user"]["full_name"], a["team"]) for a in opp["assignments"]] == [("Sales Person", "sales")]

    # Checklist: exactly the 10 verified submission documents, each cited on page 4.
    checklist = body["checklist"]
    assert sorted(i["source_clause"] for i in checklist) == sorted(d["clause_text"] for d in EXPECTED["submission_documents"])
    assert all(i["source_file_name"] == "RFP.pdf" and i["source_page"] == 4 for i in checklist)
    by_title = {i["title"]: i for i in checklist}
    assert by_title["ISO 27001 certificate"]["status"] == "available"
    assert by_title["ISO 27001 certificate"]["prefill_reason"] == "ISO 27001 recorded in the company profile"
    assert by_title["EMD Bank Guarantee or Udyam/NSIC exemption certificate"]["status"] == "available"
    assert by_title["PAN card and GST registration certificate"]["status"] == "to_prepare"
    assert sum(i["status"] == "available" for i in checklist) == 2
    assert by_title["Self-declaration of not being blacklisted"]["form_reference"] == "Annexure 3"

    # Draft: grounded, every requirement covered, skipped R7 is a placeholder.
    draft = body["latest_draft"]
    content = draft["content"]
    assert draft["version"] == 1
    assert content["disclaimer"].startswith("AI DRAFT")
    assert [t["ref"] for t in content["technical_compliance"]] == ["R1", "R4", "R5", "R6", "R7"]
    r7 = content["technical_compliance"][-1]
    assert r7["response"].startswith("[TO FILL") and r7["citation"]["page_number"] == 2
    assert [m["role"] for m in content["manpower_plan"]] == ["Java Developer", "Project Manager"]
    turnover = next(e for e in content["eligibility_compliance"] if e["criterion_type"] == "turnover")
    assert turnover["assessment"] == "MET" and turnover["citation"]["page_number"] == 3
    similar = next(e for e in content["eligibility_compliance"] if e["criterion_type"] == "similar_projects")
    assert similar["assessment"] == "UNKNOWN" and similar["statement"].startswith("[TO FILL")
    assert len(content["enclosures"]) == 10
    assert len(draft["placeholders"]) >= 6

    actions = {a.action for a in db_session.query(AuditLog)}
    assert {"bid.preparation_started", "bid.draft_generated", "opportunity.created"} <= actions


def test_the_three_downloads_open(client, sample_tender, storage_backend, llm, admin, sales):
    _analysed_tender(client, sample_tender, admin)
    opp_id = client.post(f"/api/v1/tenders/{sample_tender.id}/bid-preparation", headers=sales).json()["opportunity_id"]
    base = f"/api/v1/opportunities/{opp_id}/bid-pack"

    docx = client.get(f"{base}/draft.docx", headers=sales)
    assert docx.status_code == 200
    assert docx.headers["content-disposition"].endswith('-draft-bid-v1.docx"')
    doc = Document(io.BytesIO(docx.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "AI DRAFT — review before use" in text and "1. Covering letter" in text
    cells = {c.text for t in doc.tables for row in t.rows for c in row.cells}
    assert "RFP.pdf p.3" in cells and "Copy of valid ISO 27001 certificate." not in cells

    xlsx = client.get(f"{base}/checklist.xlsx", headers=sales)
    assert xlsx.status_code == 200
    sheet = load_workbook(io.BytesIO(xlsx.content))["Checklist"]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 10
    assert {r[6] for r in rows} == {"RFP.pdf p.4"}
    assert ("ISO 27001 certificate", "Available") in {(r[1], r[8]) for r in rows}

    archive = zipfile.ZipFile(io.BytesIO(client.get(f"/api/v1/opportunities/{opp_id}/bid-pack.zip", headers=sales).content))
    names = archive.namelist()
    assert any(n.endswith("-draft-bid.docx") for n in names) and any(n.endswith("-checklist.xlsx") for n in names)

    assert client.get(f"{base}/draft.docx").status_code == 401  # never a public link


def test_an_ungrounded_draft_is_rejected_and_not_stored(
    client, sample_tender, storage_backend, llm, admin, sales, db_session
):
    _analysed_tender(client, sample_tender, admin)
    responses = copy.deepcopy(load_rfp_responses())
    responses["bid_draft"]["company_overview"] += " We are CMMI Level 5 appraised with a turnover of ₹50 Cr."
    llm.responses = responses

    body = client.post(f"/api/v1/tenders/{sample_tender.id}/bid-preparation", headers=sales).json()

    assert body["latest_draft"] is None
    assert "rejected" in body["draft_error"]
    assert any("CMMI Level 5" in p for p in body["draft_problems"])
    assert any("₹50 Cr" in p for p in body["draft_problems"])
    assert len(body["checklist"]) == 10  # the checklist doesn't depend on the draft
    rejected = db_session.query(AuditLog).filter_by(action="bid.draft_rejected").one()
    assert len(rejected.details["problems"]) == 2

    regenerate = client.post(f"/api/v1/opportunities/{body['opportunity_id']}/bid-drafts", headers=sales)
    assert regenerate.status_code == 422 and len(regenerate.json()["detail"]["problems"]) == 2


def test_draft_needs_a_profile_and_a_configured_model(
    client, sample_tender, storage_backend, llm, admin, sales, monkeypatch
):
    _analysed_tender(client, sample_tender, admin, with_profile=False)
    body = client.post(f"/api/v1/tenders/{sample_tender.id}/bid-preparation", headers=sales).json()
    assert body["draft_error"] == "Set up the company profile before a draft can be written"
    assert len(body["checklist"]) == 10
    assert all(i["status"] == "to_prepare" for i in body["checklist"])  # nothing to confirm against

    client.put("/api/v1/profile", headers=admin, json=PROFILE)

    def unconfigured():
        raise LLMNotConfiguredError("LLM_PROVIDER=anthropic but ANTHROPIC_API_KEY is not set")

    monkeypatch.setattr("app.services.bid_preparation_service.get_llm_client", unconfigured)
    again = client.post(f"/api/v1/opportunities/{body['opportunity_id']}/bid-drafts", headers=sales)
    assert again.status_code == 503 and "ANTHROPIC_API_KEY" in again.json()["detail"]


def test_unanalysed_tender_gets_an_empty_checklist_and_a_clear_message(client, sample_tender, llm, admin, sales):
    client.put("/api/v1/profile", headers=admin, json=PROFILE)
    body = client.post(f"/api/v1/tenders/{sample_tender.id}/bid-preparation", headers=sales).json()
    assert body["checklist"] == [] and body["tender_analysed"] is False
    assert "Analyse the tender documents first" in body["draft_error"]


def test_regenerating_adds_a_version(client, sample_tender, storage_backend, llm, admin, sales):
    _analysed_tender(client, sample_tender, admin)
    opp_id = client.post(f"/api/v1/tenders/{sample_tender.id}/bid-preparation", headers=sales).json()["opportunity_id"]

    v2 = client.post(f"/api/v1/opportunities/{opp_id}/bid-drafts", headers=sales)
    assert v2.status_code == 201 and v2.json()["version"] == 2
    prep = client.get(f"/api/v1/opportunities/{opp_id}/bid-preparation", headers=sales).json()
    assert [v["version"] for v in prep["draft_versions"]] == [2, 1]
    old = client.get(f"/api/v1/opportunities/{opp_id}/bid-pack/draft.docx?version=1", headers=sales)
    assert old.headers["content-disposition"].endswith('-draft-bid-v1.docx"')


def test_checklist_edits_survive_a_reanalysis(client, sample_tender, storage_backend, llm, admin, sales):
    _analysed_tender(client, sample_tender, admin)
    body = client.post(f"/api/v1/tenders/{sample_tender.id}/bid-preparation", headers=sales).json()
    opp_id = body["opportunity_id"]
    pan = next(i for i in body["checklist"] if i["title"].startswith("PAN card"))

    patched = client.patch(
        f"/api/v1/opportunities/{opp_id}/checklist/{pan['id']}",
        headers=sales,
        json={"status": "available", "note": "Scanned copies in the shared drive"},
    ).json()
    assert patched["status"] == "available" and patched["status_set_by_user"] is True

    # Re-analysis replaces the fact rows; a sync must re-link, not duplicate or reset.
    client.post(f"/api/v1/tenders/{sample_tender.id}/analysis", headers=admin)
    synced = client.post(f"/api/v1/opportunities/{opp_id}/checklist/sync", headers=sales).json()["checklist"]
    assert len(synced) == 10
    pan_after = next(i for i in synced if i["id"] == pan["id"])
    assert pan_after["status"] == "available" and pan_after["note"] == "Scanned copies in the shared drive"


def test_manual_items_and_deletion_rules(client, sample_tender, storage_backend, llm, admin, sales):
    _analysed_tender(client, sample_tender, admin)
    body = client.post(f"/api/v1/tenders/{sample_tender.id}/bid-preparation", headers=sales).json()
    base = f"/api/v1/opportunities/{body['opportunity_id']}/checklist"

    manual = client.post(base, headers=sales, json={"title": "Board resolution"}).json()
    assert manual["is_manual"] is True and manual["position"] == 11
    assert client.delete(f"{base}/{body['checklist'][0]['id']}", headers=sales).status_code == 400
    assert client.delete(f"{base}/{manual['id']}", headers=sales).status_code == 204


def test_who_can_prepare_a_bid(client, sample_tender, storage_backend, llm, admin):
    _analysed_tender(client, sample_tender, admin)
    for role in ("viewer", "technical"):
        headers = _login(client, f"{role}-b@stackerbee.com", role)
        assert client.post(f"/api/v1/tenders/{sample_tender.id}/bid-preparation", headers=headers).status_code == 403
    bd = _login(client, "bd-b@stackerbee.com", "bd_manager")
    assert client.post(f"/api/v1/tenders/{sample_tender.id}/bid-preparation", headers=bd).status_code == 200


def test_preparing_a_bid_on_an_auto_promoted_tender_reuses_it(
    client, sample_tender, storage_backend, llm, admin, sales, monkeypatch
):
    monkeypatch.setattr(get_settings(), "opportunity_auto_promote_min_score", 1)
    _analysed_tender(client, sample_tender, admin)  # qualification runs and promotes it
    auto = client.get(f"/api/v1/tenders/{sample_tender.id}/opportunity", headers=sales).json()["opportunity"]
    assert auto["created_from"] == "auto_relevance"

    body = client.post(f"/api/v1/tenders/{sample_tender.id}/bid-preparation", headers=sales).json()
    assert body["opportunity_created"] is False and body["opportunity_id"] == auto["id"]
