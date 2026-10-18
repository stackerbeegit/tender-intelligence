"""
One-off generator for tests/fixtures/documents/*. Not run automatically —
re-run manually if the fixtures ever need regenerating:

    .venv/bin/python scripts/generate_doc_fixtures.py

Produces genuinely real files (not hand-crafted byte strings pretending to
be a PDF): a DOCX authored with python-docx, converted to a real
multi-paragraph native-text PDF via LibreOffice headless, an image-only
PDF with no text layer at all (for exercising the OCR-needed path), a
real XLSX, and a ZIP bundling two of them — matching the formats brief
§7 calls out (PDF/DOC/DOCX/XLS/XLSX/ZIP).
"""
import subprocess
import zipfile
from io import BytesIO
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pdf2image import convert_from_path
from PIL import Image

FIXTURES_DIR = Path(__file__).parent.parent / "tests" / "fixtures" / "documents"


def make_docx() -> Path:
    doc = Document()
    doc.add_heading("Sample Tender Notice", level=1)
    doc.add_paragraph(
        "This is a native-text document generated for Sprint 3 document processing tests. "
        "It contains real paragraphs so pypdf/python-docx extraction can be verified against "
        "known expected text, not a mock."
    )
    doc.add_paragraph("Minimum average annual turnover: Rs. 10,00,00,000 (Rupees Ten Crore only).")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Item"
    table.rows[0].cells[1].text = "Quantity"
    table.rows[1].cells[0].text = "Laptops"
    table.rows[1].cells[1].text = "25"

    path = FIXTURES_DIR / "sample_native.docx"
    doc.save(path)
    return path


def make_native_pdf(docx_path: Path) -> Path:
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(FIXTURES_DIR), str(docx_path)],
        check=True,
        capture_output=True,
    )
    return FIXTURES_DIR / "sample_native.pdf"


def make_scanned_pdf(native_pdf_path: Path) -> Path:
    """Rasterises the native PDF's first page into an image, then saves
    that image back out as a PDF with no text layer at all — a faithful
    stand-in for a scanned document, not a mock of one."""
    images = convert_from_path(str(native_pdf_path), dpi=150)
    first_page: Image.Image = images[0]
    path = FIXTURES_DIR / "sample_scanned.pdf"
    first_page.save(path, "PDF")
    return path


def make_xlsx() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "BOQ"
    ws.append(["Item", "Description", "Quantity", "Unit Price"])
    ws.append(["1", "Laptop - 16GB RAM", 25, 55000])
    ws.append(["2", "Networking Switch 24-port", 5, 18000])
    path = FIXTURES_DIR / "sample.xlsx"
    wb.save(path)
    return path


def make_zip(pdf_path: Path, xlsx_path: Path) -> Path:
    path = FIXTURES_DIR / "sample_bundle.zip"
    with zipfile.ZipFile(path, "w") as zf:
        zf.write(pdf_path, arcname="NIT.pdf")
        zf.write(xlsx_path, arcname="BOQ.xlsx")
    return path


RFP_PAGES: list[tuple[str, list[str]]] = [
    (
        "Request for Proposal: Development of Integrated Citizen Services Portal",
        [
            "Issued by: Department of Information Technology, Government of Punjab.",
            "RFP No: DIT/PB/2026/ICSP-014",
            "The Department invites proposals from eligible bidders for the design, development, "
            "deployment and maintenance of an Integrated Citizen Services Portal with web and mobile "
            "applications.",
            "The contract period shall be 36 months from the date of Go-Live, including 12 months of "
            "warranty support.",
        ],
    ),
    (
        "Section 4: Technical and Manpower Requirements",
        [
            "The solution shall be developed using an open-source technology stack with PostgreSQL "
            "as the database.",
            "The portal shall integrate with the State Payment Gateway and DigiLocker through "
            "secure REST APIs.",
            "The bidder shall provide 24x7 helpdesk support with a maximum resolution time of 4 hours "
            "for critical incidents.",
            "The application must undergo a security audit by a CERT-In empanelled auditor before Go-Live.",
            "Project Manager - 1 resource with minimum 10 years of experience in IT project delivery.",
            "Java Developer - 4 resources with minimum 5 years of experience each.",
        ],
    ),
    (
        "Section 5: Pre-Qualification Criteria",
        [
            "The bidder must have an average annual turnover of at least Rs. 10 Crore during the last "
            "three financial years.",
            "The bidder must have successfully completed at least 3 similar projects, each of value not "
            "less than Rs. 2 Crore, for Government or PSU clients.",
            "The bidder must hold a valid ISO 27001 certification.",
            "Earnest Money Deposit of Rs. 5,00,000 shall be submitted along with the bid.",
            "Micro and Small Enterprises registered with NSIC or Udyam are exempted from payment of "
            "EMD as per the Public Procurement Policy for MSEs Order, 2012.",
        ],
    ),
]


def make_rfp_pdf() -> Path:
    """Sprint 4: a representative 3-page RFP whose clauses
    tests/fixtures/intelligence/sample_rfp_expected.json describes. One
    section per page (explicit page breaks) so page citations are
    deterministic."""
    doc = Document()
    for index, (heading, paragraphs) in enumerate(RFP_PAGES):
        if index:
            doc.add_page_break()
        doc.add_heading(heading, level=1)
        for text in paragraphs:
            doc.add_paragraph(text)

    docx_path = FIXTURES_DIR / "sample_rfp.docx"
    doc.save(docx_path)
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(FIXTURES_DIR), str(docx_path)],
        check=True,
        capture_output=True,
    )
    docx_path.unlink()  # only the PDF is a fixture
    return FIXTURES_DIR / "sample_rfp.pdf"


def main() -> None:
    FIXTURES_DIR.mkdir(parents=True, exist_ok=True)
    docx_path = make_docx()
    pdf_path = make_native_pdf(docx_path)
    scanned_path = make_scanned_pdf(pdf_path)
    xlsx_path = make_xlsx()
    zip_path = make_zip(pdf_path, xlsx_path)
    rfp_path = make_rfp_pdf()

    for p in (docx_path, pdf_path, scanned_path, xlsx_path, zip_path, rfp_path):
        print(f"{p.name}: {p.stat().st_size} bytes")


if __name__ == "__main__":
    main()
